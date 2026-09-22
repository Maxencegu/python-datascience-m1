"""
Tableau de bord enseignant — suivi complet de la promo.

Usage :
    python templates/dashboard_etudiants.py                 # tous les TD
    python templates/dashboard_etudiants.py --td td03       # un seul TD
    python templates/dashboard_etudiants.py --refresh       # ignore le cache GitHub
    python templates/dashboard_etudiants.py --offline       # aucune requête GitHub

Pour chaque étudiant et chaque TD, le script collecte :
  - la présence (export Google Forms)
  - le résultat du quiz (export Google Forms)
  - la présence du notebook tdXX_enonce.ipynb sur main
  - la comparaison des sorties de cellules avec templates/tests/tdXX_expected.json
  - le cycle Git : branche dev_tdXX, Pull Request, review par un tiers, merge,
    suppression de la branche, conclusion des GitHub Actions
  - les dates de création et de merge de la Pull Request
puis calcule les notes selon le barème du README et écrit un tableau de bord
HTML autonome.

Données d'entrée — toutes dans réponses/ (dossier ignoré par git, jamais publié) :
    TD01 - Feuille de Présence.csv(.zip|.xlsx)   export Google Forms
    TD02 - Quiz.csv(.zip|.xlsx)                  export Google Forms
    etudiants.csv                                cache e-mail → pseudo GitHub,
                                                 complété automatiquement depuis
                                                 l'URL du dépôt saisie au quiz,
                                                 éditable à la main

Sortie : réponses/dashboard.html (contient des données nominatives — ne jamais committer)

Prérequis : gh CLI authentifié (gh auth login). openpyxl seulement si vos
exports sont au format .xlsx.
"""
import argparse
import csv
import io
import json
import os
import re
import subprocess
import sys
import time
import unicodedata
import zipfile
from datetime import datetime, timezone

# ── Emplacements ────────────────────────────────────────────────────────────
RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REPONSES = os.path.join(RACINE, "réponses")
TESTS = os.path.join(RACINE, "templates", "tests")
CACHE = os.path.join(REPONSES, ".cache_github.json")
MAPPING = os.path.join(REPONSES, "etudiants.csv")
SORTIE = os.path.join(REPONSES, "dashboard.html")

REPO_NAME = "upjv-python-datascience"
CACHE_TTL = 15 * 60  # secondes

# ── Barème (README.md § Évaluation) ─────────────────────────────────────────
TDS = ["td01", "td02", "td03", "td04", "td05", "td06", "td07", "td08"]

LIBELLES = {
    "td01": "Introduction à Git",
    "td02": "Introduction à GitHub",
    "td03": "Les bases de Python",
    "td04": "Fonctions, modules et NumPy",
    "td05": "Matplotlib, dictionnaires et Pandas",
    "td06": "Logique, flux de contrôle et boucles",
    "td07": "Étude de cas + mini-projet 1",
    "td08": "Pandas : agrégation et transformation",
}

# points par TD : présence / notebook / quiz / mini-projet
BAREME = {td: {"presence": 10, "notebook": 20, "quiz": 10, "mp": 0} for td in TDS}
BAREME["td01"]["quiz"] = 0                      # quiz à partir du TD02
BAREME["td07"] = {"presence": 10, "notebook": 20, "quiz": 0, "mp": 30}

# Le cycle Git (branche → PR → review → merge) démarre au TD02.
PREMIER_TD_GIT = "td02"


# ════════════════════════════════════════════════════════════════════════════
#  Lecture des exports Google Forms
# ════════════════════════════════════════════════════════════════════════════
def _normaliser(texte):
    """minuscules, sans accents ni ponctuation — pour comparer des en-têtes."""
    texte = unicodedata.normalize("NFD", str(texte))
    texte = "".join(c for c in texte if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9]+", " ", texte.lower()).strip()


def _lire_csv(contenu):
    """Décode et parse un CSV Google Forms (utf-8 ou latin-1, séparateur variable)."""
    for encodage in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            texte = contenu.decode(encodage)
            break
        except UnicodeDecodeError:
            continue
    else:
        return []
    try:
        dialecte = csv.Sniffer().sniff(texte[:4096], delimiters=",;\t")
    except csv.Error:
        dialecte = csv.excel
    return list(csv.DictReader(io.StringIO(texte), dialect=dialecte))


def lire_tableur(chemin):
    """Retourne une liste de dict pour un .csv, un .csv.zip ou un .xlsx."""
    if chemin.endswith(".zip"):
        with zipfile.ZipFile(chemin) as z:
            noms = [n for n in z.namelist() if n.lower().endswith(".csv")]
            return _lire_csv(z.read(noms[0])) if noms else []
    if chemin.endswith(".xlsx"):
        try:
            import openpyxl
        except ImportError:
            print(f"  ⚠️  openpyxl absent — {os.path.basename(chemin)} ignoré "
                  f"(pip install openpyxl, ou exportez en CSV)")
            return []
        classeur = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
        feuille = classeur[classeur.sheetnames[0]]
        lignes = feuille.iter_rows(values_only=True)
        entetes = [str(c) if c is not None else "" for c in next(lignes, [])]
        return [dict(zip(entetes, ["" if c is None else str(c) for c in ligne]))
                for ligne in lignes]
    with open(chemin, "rb") as f:
        return _lire_csv(f.read())


def colonne(ligne, *mots):
    """Valeur de la première colonne dont l'en-tête contient tous les mots."""
    for cle, valeur in ligne.items():
        entete = _normaliser(cle)
        if all(_normaliser(m) in entete for m in mots):
            return (valeur or "").strip()
    return ""


def fichiers_exports(motif):
    """Fichiers de réponses/ dont le nom contient `motif`, groupés par TD."""
    trouves = {}
    if not os.path.isdir(REPONSES):
        return trouves
    for nom in sorted(os.listdir(REPONSES)):
        if not nom.lower().endswith((".csv", ".zip", ".xlsx")):
            continue
        normalise = _normaliser(nom)
        if _normaliser(motif) not in normalise:
            continue
        correspondance = re.search(r"td\s*0*(\d+)", normalise)
        if correspondance:
            trouves.setdefault(f"td{int(correspondance.group(1)):02d}", []).append(
                os.path.join(REPONSES, nom))
    return trouves


def charger_presences():
    """{td: {email: horodatage}}"""
    presences = {}
    for td, chemins in fichiers_exports("presence").items():
        presences[td] = {}
        for chemin in chemins:
            for ligne in lire_tableur(chemin):
                email = colonne(ligne, "utilisateur") or colonne(ligne, "adresse", "mail")
                if email:
                    presences[td][email.strip().lower()] = colonne(ligne, "horodateur")
    return presences


def _score_quiz(ligne):
    """Google Forms écrit le score sous la forme '8 / 10'."""
    brut = colonne(ligne, "score") or colonne(ligne, "note")
    correspondance = re.search(r"([\d.,]+)\s*/\s*([\d.,]+)", brut)
    if correspondance:
        obtenu = float(correspondance.group(1).replace(",", "."))
        total = float(correspondance.group(2).replace(",", "."))
        return obtenu, total
    if brut.replace(",", ".").replace(".", "").isdigit():
        return float(brut.replace(",", ".")), None
    return None, None


def charger_quiz():
    """{td: {email: {"obtenu":…, "total":…, "url":…}}}"""
    quiz = {}
    for td, chemins in fichiers_exports("quiz").items():
        quiz[td] = {}
        for chemin in chemins:
            for ligne in lire_tableur(chemin):
                email = colonne(ligne, "utilisateur") or colonne(ligne, "adresse", "mail")
                if not email:
                    continue
                obtenu, total = _score_quiz(ligne)
                url = (colonne(ligne, "url", "depot") or colonne(ligne, "lien", "depot")
                       or colonne(ligne, "url") or colonne(ligne, "github"))
                quiz[td][email.strip().lower()] = {
                    "obtenu": obtenu, "total": total, "url": url,
                    "horodatage": colonne(ligne, "horodateur"),
                }
    return quiz


# ════════════════════════════════════════════════════════════════════════════
#  Annuaire des étudiants
# ════════════════════════════════════════════════════════════════════════════
CHAMPS_MAPPING = ["email", "nom", "numero", "username"]


def username_depuis_url(url):
    correspondance = re.search(r"github\.com/([A-Za-z0-9](?:[A-Za-z0-9-]*[A-Za-z0-9])?)", url or "")
    return correspondance.group(1) if correspondance else ""


def charger_mapping():
    annuaire = {}
    if os.path.exists(MAPPING):
        for ligne in lire_tableur(MAPPING):
            email = (ligne.get("email") or "").strip().lower()
            if email:
                annuaire[email] = {c: (ligne.get(c) or "").strip() for c in CHAMPS_MAPPING}
                annuaire[email]["email"] = email
    return annuaire


def sauver_mapping(annuaire):
    os.makedirs(REPONSES, exist_ok=True)
    with open(MAPPING, "w", encoding="utf-8", newline="") as f:
        redacteur = csv.DictWriter(f, fieldnames=CHAMPS_MAPPING)
        redacteur.writeheader()
        for email in sorted(annuaire):
            redacteur.writerow(annuaire[email])


def construire_annuaire(presences, quiz):
    """Fusionne le cache, les feuilles de présence et les URL saisies au quiz."""
    annuaire = charger_mapping()

    for td, chemins in fichiers_exports("presence").items():
        for chemin in chemins:
            for ligne in lire_tableur(chemin):
                email = (colonne(ligne, "utilisateur") or colonne(ligne, "adresse", "mail")).strip().lower()
                if not email:
                    continue
                fiche = annuaire.setdefault(email, {c: "" for c in CHAMPS_MAPPING})
                fiche["email"] = email
                fiche["nom"] = fiche["nom"] or colonne(ligne, "nom", "prenom") or colonne(ligne, "nom")
                fiche["numero"] = fiche["numero"] or colonne(ligne, "numero", "etudiant")

    for td in sorted(quiz):
        for email, reponse in quiz[td].items():
            fiche = annuaire.setdefault(email, {c: "" for c in CHAMPS_MAPPING})
            fiche["email"] = email
            username = username_depuis_url(reponse.get("url", ""))
            if username and not fiche.get("username"):
                fiche["username"] = username

    sauver_mapping(annuaire)
    return annuaire


# ════════════════════════════════════════════════════════════════════════════
#  Accès GitHub (gh CLI + cache disque)
# ════════════════════════════════════════════════════════════════════════════
class GitHub:
    def __init__(self, hors_ligne=False, rafraichir=False):
        self.hors_ligne = hors_ligne
        self.cache = {}
        if os.path.exists(CACHE) and not rafraichir:
            try:
                with open(CACHE, encoding="utf-8") as f:
                    self.cache = json.load(f)
            except json.JSONDecodeError:
                self.cache = {}
        self.appels = 0

    def enregistrer(self):
        os.makedirs(REPONSES, exist_ok=True)
        with open(CACHE, "w", encoding="utf-8") as f:
            json.dump(self.cache, f)

    def api(self, chemin, paginer=False):
        """GET sur l'API GitHub. Retourne l'objet décodé, ou None si absent."""
        entree = self.cache.get(chemin)
        if entree and time.time() - entree["t"] < CACHE_TTL:
            return entree["v"]
        if self.hors_ligne:
            return entree["v"] if entree else None

        commande = ["gh", "api", chemin]
        if paginer:
            commande.append("--paginate")
        resultat = subprocess.run(commande, capture_output=True, text=True, encoding="utf-8")
        self.appels += 1
        if resultat.returncode != 0:
            valeur = None
        else:
            try:
                valeur = json.loads(resultat.stdout or "null")
            except json.JSONDecodeError:
                valeur = None
        self.cache[chemin] = {"t": time.time(), "v": valeur}
        return valeur

    def fichier(self, depot, chemin, ref="main"):
        """Contenu texte d'un fichier du dépôt, ou None."""
        cle = f"RAW:{depot}/{ref}/{chemin}"
        entree = self.cache.get(cle)
        if entree and time.time() - entree["t"] < CACHE_TTL:
            return entree["v"]
        if self.hors_ligne:
            return entree["v"] if entree else None

        resultat = subprocess.run(
            ["gh", "api", f"repos/{depot}/contents/{chemin}?ref={ref}",
             "--header", "Accept: application/vnd.github.raw"],
            capture_output=True, text=True, encoding="utf-8",
        )
        self.appels += 1
        valeur = resultat.stdout if resultat.returncode == 0 else None
        self.cache[cle] = {"t": time.time(), "v": valeur}
        return valeur


# ════════════════════════════════════════════════════════════════════════════
#  Audit des sorties d'un notebook
# ════════════════════════════════════════════════════════════════════════════
SIGNES_ECHEC = ("❌", "⏳", "💡", "⚠️")
ANSI = re.compile(r"\x1b\[[0-9;]*m")


def attendus(td):
    chemin = os.path.join(TESTS, f"{td}_expected.json")
    if os.path.exists(chemin):
        with open(chemin, encoding="utf-8") as f:
            return json.load(f)
    return None


def sortie_cellule(cellule):
    morceaux = []
    for sortie in cellule.get("outputs", []):
        if sortie.get("output_type") == "stream":
            morceaux.append("".join(sortie.get("text", [])))
        elif sortie.get("output_type") in ("execute_result", "display_data"):
            texte = "".join(sortie.get("data", {}).get("text/plain", []))
            if not texte.startswith(("<IPython.core.display.", "<IPython.lib.display.")):
                morceaux.append(texte)
    return ANSI.sub("", "".join(morceaux)).strip()


def erreur_cellule(cellule):
    for sortie in cellule.get("outputs", []):
        if sortie.get("output_type") == "error":
            return f"{sortie.get('ename')}: {sortie.get('evalue')}"
    return None


def auditer_notebook(source, td):
    """Compare les sorties enregistrées aux attendus (ou, à défaut, cherche les ❌)."""
    resultat = {"present": False, "ok": 0, "total": 0, "details": [], "placeholders": 0}
    if not source:
        return resultat
    try:
        notebook = json.loads(source)
    except json.JSONDecodeError:
        resultat["details"].append("notebook illisible (JSON invalide)")
        return resultat

    resultat["present"] = True
    cellules = notebook.get("cells", [])
    reference = attendus(td)

    if reference:
        resultat["total"] = len(reference)
        for index, attendu in sorted(reference.items(), key=lambda kv: int(kv[0])):
            i = int(index)
            if i >= len(cellules) or cellules[i].get("cell_type") != "code":
                resultat["details"].append(f"cellule {i} absente ou déplacée")
                continue
            obtenu = sortie_cellule(cellules[i])
            if obtenu == attendu.strip():
                resultat["ok"] += 1
            else:
                apercu = (obtenu or "aucune sortie")[:60].replace("\n", " ⏎ ")
                resultat["details"].append(f"cellule {i} : attendu {attendu.strip()[:40]!r}, obtenu {apercu!r}")
    else:
        # TD sans fichier d'attendus (TD01, TD02) : on lit les cellules « ✅ Vérification ».
        for i, cellule in enumerate(cellules):
            if cellule.get("cell_type") != "code":
                continue
            source_cellule = "".join(cellule.get("source", []))
            premiere_ligne = source_cellule.split("\n")[0]
            if "Vérification" not in premiere_ligne:
                continue
            # La vérification du cycle Pull Request s'exécute après l'upload :
            # elle ne peut pas être verte dans le fichier déposé. Évaluée par l'API.
            if f"dev_{td}" in source_cellule:
                continue
            resultat["total"] += 1
            texte = sortie_cellule(cellule)
            if not texte:
                resultat["details"].append(f"cellule {i} : jamais exécutée")
            elif any(signe in texte for signe in SIGNES_ECHEC):
                ligne = next(l for l in texte.split("\n") if any(s in l for s in SIGNES_ECHEC))
                resultat["details"].append(f"cellule {i} : {ligne.strip()[:70]}")
            else:
                resultat["ok"] += 1

    for i, cellule in enumerate(cellules):
        if cellule.get("cell_type") == "code":
            if "____" in "".join(cellule.get("source", [])):
                resultat["placeholders"] += 1
            erreur = erreur_cellule(cellule)
            if erreur:
                resultat["details"].append(f"cellule {i} : {erreur}")
    return resultat


# ════════════════════════════════════════════════════════════════════════════
#  Cycle Git d'un TD
# ════════════════════════════════════════════════════════════════════════════
def cycle_git(gh, depot, proprietaire, td, branches):
    """Branche, Pull Request, review par un tiers, merge, nettoyage, Actions."""
    branche = f"dev_{td}"
    etat = {
        "branche_existe": branche in branches,
        "pr": False, "pr_numero": None, "creee_le": None, "mergee_le": None,
        "approuvee_par": None, "auto_approuvee": False,
        "branche_supprimee": False, "checks": None, "url": None,
    }
    prs = gh.api(f"repos/{depot}/pulls?state=all&base=main&head={proprietaire}:{branche}") or []
    if not prs:
        etat["branche_supprimee"] = not etat["branche_existe"] and False
        return etat

    pr = next((p for p in prs if p.get("merged_at")), prs[0])
    etat.update({
        "pr": True,
        "pr_numero": pr.get("number"),
        "creee_le": pr.get("created_at"),
        "mergee_le": pr.get("merged_at"),
        "url": pr.get("html_url"),
        "branche_supprimee": not etat["branche_existe"],
    })

    revues = gh.api(f"repos/{depot}/pulls/{pr['number']}/reviews") or []
    approbations = [r for r in revues if r.get("state") == "APPROVED"]
    for revue in approbations:
        auteur = (revue.get("user") or {}).get("login", "")
        if auteur and auteur.lower() != proprietaire.lower():
            etat["approuvee_par"] = auteur
            break
    else:
        if approbations:
            etat["auto_approuvee"] = True

    sha = (pr.get("head") or {}).get("sha")
    if sha:
        runs = gh.api(f"repos/{depot}/commits/{sha}/check-runs") or {}
        conclusions = [c.get("conclusion") for c in runs.get("check_runs", [])]
        if conclusions:
            etat["checks"] = "success" if all(c == "success" for c in conclusions) else "failure"
    return etat


# ════════════════════════════════════════════════════════════════════════════
#  Collecte + notes
# ════════════════════════════════════════════════════════════════════════════
def notes_du_td(td, presence, quiz, audit, git):
    bareme = BAREME[td]
    notes = {"presence": 0.0, "notebook": 0.0, "quiz": 0.0, "mp": 0.0}

    if presence:
        notes["presence"] = float(bareme["presence"])

    if audit["total"]:
        notes["notebook"] = round(bareme["notebook"] * audit["ok"] / audit["total"], 1)

    if bareme["quiz"] and quiz and quiz.get("obtenu") is not None:
        total = quiz.get("total") or bareme["quiz"]
        notes["quiz"] = round(bareme["quiz"] * quiz["obtenu"] / total, 1)

    obtenu = sum(notes.values())
    maximum = sum(bareme.values())
    return {"detail": notes, "obtenu": round(obtenu, 1), "max": maximum}


def collecter(tds, hors_ligne=False, rafraichir=False):
    presences = charger_presences()
    quiz = charger_quiz()
    annuaire = construire_annuaire(presences, quiz)
    gh = GitHub(hors_ligne=hors_ligne, rafraichir=rafraichir)

    etudiants = []
    for email in sorted(annuaire, key=lambda e: (annuaire[e].get("nom") or e).lower()):
        fiche = annuaire[email]
        username = fiche.get("username", "")
        depot = f"{username}/{REPO_NAME}" if username else None

        branches = []
        depot_ok = False
        if depot:
            infos = gh.api(f"repos/{depot}")
            depot_ok = bool(infos)
            if depot_ok:
                branches = [b["name"] for b in (gh.api(f"repos/{depot}/branches?per_page=100") or [])]

        donnees_tds = {}
        for td in tds:
            audit = {"present": False, "ok": 0, "total": 0, "details": [], "placeholders": 0}
            git = None
            if depot_ok:
                audit = auditer_notebook(gh.fichier(depot, f"{td}_enonce.ipynb"), td)
                if td >= PREMIER_TD_GIT:
                    git = cycle_git(gh, depot, username, td, branches)

            presence = email in presences.get(td, {})
            note = notes_du_td(td, presence, quiz.get(td, {}).get(email), audit, git)
            donnees_tds[td] = {
                "presence": presence,
                "presence_le": presences.get(td, {}).get(email),
                "quiz": quiz.get(td, {}).get(email),
                "audit": audit,
                "git": git,
                "note": note,
            }

        total = round(sum(t["note"]["obtenu"] for t in donnees_tds.values()), 1)
        maximum = sum(t["note"]["max"] for t in donnees_tds.values())
        etudiants.append({
            "nom": fiche.get("nom") or email,
            "email": email,
            "numero": fiche.get("numero", ""),
            "username": username,
            "depot": f"https://github.com/{depot}" if depot else "",
            "depot_ok": depot_ok,
            "tds": donnees_tds,
            "total": total,
            "max": maximum,
            "pourcentage": round(100 * total / maximum, 1) if maximum else 0,
        })

    gh.enregistrer()
    return {
        "genere_le": datetime.now(timezone.utc).astimezone().strftime("%d/%m/%Y %H:%M"),
        "tds": tds,
        "libelles": LIBELLES,
        "bareme": BAREME,
        "etudiants": etudiants,
        "appels_api": gh.appels,
    }


# ════════════════════════════════════════════════════════════════════════════
#  Rendu HTML
# ════════════════════════════════════════════════════════════════════════════
GABARIT = """<!DOCTYPE html>
<html lang="fr">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Tableau de bord — Python &amp; Data Science</title>
<style>
  :root {
    --fond: #0d1117; --carte: #161b22; --bordure: #21262d; --bordure-claire: #30363d;
    --texte: #e6edf3; --doux: #8b949e; --bleu: #58a6ff; --violet: #8b5cf6;
    --vert: #3fb950; --orange: #d29922; --rouge: #f85149;
  }
  * { box-sizing: border-box; margin: 0; padding: 0; }
  body {
    background: var(--fond); color: var(--texte); padding: 28px 22px 60px;
    font-family: "Nunito", "Segoe UI", system-ui, sans-serif;
  }
  .entete { text-align: center; margin-bottom: 26px; }
  .fil {
    color: var(--bleu); font-size: 13px; font-weight: 800; letter-spacing: .15em;
    text-transform: uppercase; margin-bottom: 8px;
  }
  h1 {
    font-size: 30px; font-weight: 800;
    background: linear-gradient(90deg, var(--bleu), var(--violet));
    -webkit-background-clip: text; -webkit-text-fill-color: transparent;
  }
  .sous { color: var(--doux); font-size: 14px; margin-top: 6px; }
  .cartes { display: flex; gap: 14px; flex-wrap: wrap; justify-content: center; margin: 22px 0; }
  .kpi {
    background: var(--carte); border: 1px solid var(--bordure); border-top: 3px solid var(--bleu);
    border-radius: 11px; padding: 13px 20px; min-width: 150px; text-align: center;
  }
  .kpi .v { font-size: 25px; font-weight: 800; }
  .kpi .l { color: var(--doux); font-size: 12.5px; text-transform: uppercase; letter-spacing: .1em; margin-top: 3px; }
  .filtres {
    display: flex; gap: 11px; flex-wrap: wrap; align-items: center;
    background: var(--carte); border: 1px solid var(--bordure); border-radius: 11px;
    padding: 13px 16px; margin-bottom: 16px;
  }
  input, select, button {
    background: var(--fond); color: var(--texte); border: 1px solid var(--bordure-claire);
    border-radius: 7px; padding: 7px 11px; font-size: 14px; font-family: inherit;
  }
  input { min-width: 230px; }
  button { cursor: pointer; }
  button.actif { border-color: var(--bleu); color: var(--bleu); }
  table { width: 100%; border-collapse: collapse; background: var(--carte);
          border: 1px solid var(--bordure); border-radius: 11px; overflow: hidden; }
  th, td { padding: 9px 12px; text-align: left; border-bottom: 1px solid var(--bordure); font-size: 14px; }
  th { background: var(--fond); font-weight: 800; cursor: pointer; white-space: nowrap; user-select: none; }
  th:hover { color: var(--bleu); }
  tr.ligne:hover { background: #1a2029; }
  tr.detail td { background: #10151c; }
  .nom { font-weight: 700; }
  .meta { color: var(--doux); font-size: 12px; }
  .pastille { display: inline-block; padding: 1px 8px; border-radius: 20px; font-size: 12.5px; font-weight: 700; }
  .ok { background: #10301b; color: #56d364; }
  .ko { background: #2a1113; color: #ff7b72; }
  .attente { background: #2b2410; color: #e3b341; }
  .neutre { background: #1b2027; color: var(--doux); }
  .barre { height: 7px; border-radius: 4px; background: #21262d; overflow: hidden; min-width: 90px; }
  .barre span { display: block; height: 100%; background: linear-gradient(90deg, var(--vert), var(--bleu)); }
  .grille { display: grid; grid-template-columns: repeat(auto-fill, minmax(290px, 1fr)); gap: 12px; padding: 6px 0 10px; }
  .fiche { background: var(--carte); border: 1px solid var(--bordure); border-radius: 10px; padding: 12px 14px; }
  .fiche h4 { font-size: 14.5px; margin-bottom: 7px; }
  .fiche .l { display: flex; justify-content: space-between; gap: 10px; font-size: 13px; padding: 2px 0; color: #c3cbd4; }
  .pb { color: #ff9c92; font-size: 12.5px; margin-top: 6px; line-height: 1.45; }
  a { color: var(--bleu); text-decoration: none; }
  .vide { text-align: center; color: var(--doux); padding: 40px; }
</style>
</head>
<body>
  <div class="entete">
    <div class="fil">Python Data Science — UPJV Amiens</div>
    <h1>Tableau de bord de la promo</h1>
    <div class="sous">Généré le __DATE__ · __APPELS__ appels API · données nominatives, ne pas publier</div>
  </div>

  <div class="cartes" id="kpis"></div>

  <div class="filtres">
    <input id="recherche" placeholder="Rechercher un étudiant, un pseudo…">
    <select id="selectTd"></select>
    <button id="btnProblemes">Problèmes uniquement</button>
    <button id="btnSansDepot">Sans dépôt</button>
    <span class="meta" id="compteur"></span>
  </div>

  <table>
    <thead id="entetes"></thead>
    <tbody id="corps"></tbody>
  </table>

<script>
const DONNEES = __DONNEES__;
let triCle = "nom", triSens = 1, filtreProblemes = false, filtreSansDepot = false;

const pastille = (etat, texte) => `<span class="pastille ${etat}">${texte}</span>`;
const jour = (iso) => iso ? new Date(iso).toLocaleDateString("fr-FR") : "—";

function tdsAffiches() {
  const choix = document.getElementById("selectTd").value;
  return choix === "tous" ? DONNEES.tds : [choix];
}

function scoreEtudiant(e) {
  const tds = tdsAffiches();
  const obtenu = tds.reduce((s, td) => s + e.tds[td].note.obtenu, 0);
  const max = tds.reduce((s, td) => s + e.tds[td].note.max, 0);
  return { obtenu: Math.round(obtenu * 10) / 10, max, pct: max ? Math.round(1000 * obtenu / max) / 10 : 0 };
}

function problemes(e) {
  const liste = [];
  if (!e.username) liste.push("pseudo GitHub inconnu");
  else if (!e.depot_ok) liste.push("dépôt introuvable ou privé");
  for (const td of tdsAffiches()) {
    const d = e.tds[td], b = DONNEES.bareme[td];
    if (!d.presence) liste.push(`${td} : absent`);
    if (e.depot_ok && !d.audit.present) liste.push(`${td} : notebook absent de main`);
    else if (d.audit.total && d.audit.ok < d.audit.total)
      liste.push(`${td} : ${d.audit.total - d.audit.ok}/${d.audit.total} cellules incorrectes`);
    if (d.audit.placeholders) liste.push(`${td} : ${d.audit.placeholders} placeholder(s) ____`);
    if (b.quiz && (!d.quiz || d.quiz.obtenu === null)) liste.push(`${td} : quiz non fait`);
    if (d.git) {
      if (!d.git.pr) liste.push(`${td} : aucune Pull Request`);
      else {
        if (!d.git.mergee_le) liste.push(`${td} : PR non mergée`);
        if (!d.git.approuvee_par) liste.push(`${td} : ${d.git.auto_approuvee ? "auto-approuvée" : "sans review"}`);
        if (!d.git.branche_supprimee) liste.push(`${td} : branche non supprimée`);
        if (d.git.checks === "failure") liste.push(`${td} : Actions en échec`);
      }
    }
  }
  return liste;
}

function etatGit(g) {
  if (!g) return pastille("neutre", "—");
  if (!g.pr) return pastille("ko", "pas de PR");
  if (!g.mergee_le) return pastille("attente", "PR ouverte");
  if (!g.approuvee_par) return pastille("attente", g.auto_approuvee ? "auto-approuvée" : "sans review");
  if (!g.branche_supprimee) return pastille("attente", "branche restante");
  if (g.checks === "failure") return pastille("ko", "Actions ❌");
  return pastille("ok", "cycle complet");
}

function lignes() {
  const recherche = document.getElementById("recherche").value.toLowerCase();
  return DONNEES.etudiants.filter(e => {
    if (recherche && !(`${e.nom} ${e.email} ${e.username} ${e.numero}`.toLowerCase().includes(recherche))) return false;
    if (filtreSansDepot && e.depot_ok) return false;
    if (filtreProblemes && problemes(e).length === 0) return false;
    return true;
  }).sort((a, b) => {
    let va, vb;
    if (triCle === "note") { va = scoreEtudiant(a).pct; vb = scoreEtudiant(b).pct; }
    else if (triCle === "problemes") { va = problemes(a).length; vb = problemes(b).length; }
    else { va = (a[triCle] || "").toString().toLowerCase(); vb = (b[triCle] || "").toString().toLowerCase(); }
    return va < vb ? -triSens : va > vb ? triSens : 0;
  });
}

function detail(e) {
  const cartes = tdsAffiches().map(td => {
    const d = e.tds[td], b = DONNEES.bareme[td], n = d.note.detail;
    const pbs = d.audit.details.slice(0, 4).map(x => `• ${x}`).join("<br>");
    return `<div class="fiche">
      <h4>${td.toUpperCase()} — ${DONNEES.libelles[td]}</h4>
      <div class="l"><span>Présence</span><span>${d.presence ? "✅" : "❌"} ${n.presence}/${b.presence}</span></div>
      <div class="l"><span>Notebook (${d.audit.ok}/${d.audit.total} cellules)</span><span>${n.notebook}/${b.notebook}</span></div>
      ${b.quiz ? `<div class="l"><span>Quiz</span><span>${n.quiz}/${b.quiz}</span></div>` : ""}
      ${b.mp ? `<div class="l"><span>Mini-projet</span><span>${n.mp}/${b.mp}</span></div>` : ""}
      ${d.git ? `<div class="l"><span>Cycle Git</span><span>${etatGit(d.git)}</span></div>
      <div class="l"><span>PR créée / mergée</span><span>${jour(d.git.creee_le)} → ${jour(d.git.mergee_le)}</span></div>
      ${d.git.approuvee_par ? `<div class="l"><span>Approuvée par</span><span>${d.git.approuvee_par}</span></div>` : ""}
      ${d.git.url ? `<div class="l"><span>Lien</span><a href="${d.git.url}" target="_blank">PR #${d.git.pr_numero}</a></div>` : ""}` : ""}
      ${pbs ? `<div class="pb">${pbs}</div>` : ""}
    </div>`;
  }).join("");
  return `<tr class="detail"><td colspan="7"><div class="grille">${cartes}</div></td></tr>`;
}

function rendre() {
  const tds = tdsAffiches();
  const donnees = lignes();
  document.getElementById("entetes").innerHTML = `<tr>
    <th data-cle="nom">Étudiant</th>
    <th data-cle="username">Pseudo GitHub</th>
    <th>Présence</th>
    <th>Notebooks</th>
    <th>Cycle Git</th>
    <th data-cle="note">Note</th>
    <th data-cle="problemes">Points bloquants</th></tr>`;

  document.getElementById("corps").innerHTML = donnees.length === 0
    ? `<tr><td colspan="7" class="vide">Aucun étudiant ne correspond au filtre.</td></tr>`
    : donnees.map((e, i) => {
      const s = scoreEtudiant(e), pbs = problemes(e);
      const presents = tds.filter(td => e.tds[td].presence).length;
      const cellulesOk = tds.reduce((a, td) => a + e.tds[td].audit.ok, 0);
      const cellulesTot = tds.reduce((a, td) => a + e.tds[td].audit.total, 0);
      const gits = tds.filter(td => e.tds[td].git).map(td => etatGit(e.tds[td].git)).join(" ");
      return `<tr class="ligne" onclick="basculer(${i})">
        <td><div class="nom">${e.nom}</div><div class="meta">${e.email}${e.numero ? " · " + e.numero : ""}</div></td>
        <td>${e.username ? `<a href="${e.depot}" target="_blank">${e.username}</a>` : pastille("ko", "inconnu")}</td>
        <td>${presents}/${tds.length}</td>
        <td>${cellulesOk}/${cellulesTot}<div class="barre"><span style="width:${cellulesTot ? 100 * cellulesOk / cellulesTot : 0}%"></span></div></td>
        <td>${gits || pastille("neutre", "—")}</td>
        <td><b>${s.obtenu}</b>/${s.max}<div class="meta">${s.pct} %</div></td>
        <td>${pbs.length ? pastille("ko", pbs.length + " à voir") : pastille("ok", "tout est propre")}</td>
      </tr>` + `<tr class="detail" id="d${i}" style="display:none"><td colspan="7"></td></tr>`;
    }).join("");

  donnees.forEach((e, i) => {
    const ligne = document.getElementById("d" + i);
    if (ligne) ligne.innerHTML = detail(e).replace(/^<tr class="detail">|<\\/tr>$/g, "");
  });

  const moyenne = donnees.length
    ? Math.round(10 * donnees.reduce((a, e) => a + scoreEtudiant(e).pct, 0) / donnees.length) / 10 : 0;
  const complets = donnees.filter(e => problemes(e).length === 0).length;
  document.getElementById("kpis").innerHTML = `
    <div class="kpi"><div class="v">${donnees.length}</div><div class="l">Étudiants</div></div>
    <div class="kpi"><div class="v">${complets}</div><div class="l">Dossiers complets</div></div>
    <div class="kpi"><div class="v">${moyenne} %</div><div class="l">Moyenne</div></div>
    <div class="kpi"><div class="v">${donnees.filter(e => !e.depot_ok).length}</div><div class="l">Sans dépôt</div></div>`;
  document.getElementById("compteur").textContent = `${donnees.length} ligne(s)`;
}

function basculer(i) {
  const ligne = document.getElementById("d" + i);
  ligne.style.display = ligne.style.display === "none" ? "" : "none";
}

document.getElementById("selectTd").innerHTML =
  `<option value="tous">Tous les TD</option>` +
  DONNEES.tds.map(td => `<option value="${td}">${td.toUpperCase()} — ${DONNEES.libelles[td]}</option>`).join("");
document.getElementById("recherche").addEventListener("input", rendre);
document.getElementById("selectTd").addEventListener("change", rendre);
document.getElementById("btnProblemes").addEventListener("click", (ev) => {
  filtreProblemes = !filtreProblemes; ev.target.classList.toggle("actif", filtreProblemes); rendre();
});
document.getElementById("btnSansDepot").addEventListener("click", (ev) => {
  filtreSansDepot = !filtreSansDepot; ev.target.classList.toggle("actif", filtreSansDepot); rendre();
});
document.addEventListener("click", (ev) => {
  const th = ev.target.closest("th[data-cle]");
  if (!th) return;
  const cle = th.dataset.cle;
  triSens = cle === triCle ? -triSens : 1; triCle = cle; rendre();
});
rendre();
</script>
</body>
</html>"""


def ecrire_html(donnees, chemin):
    html = (GABARIT
            .replace("__DATE__", donnees["genere_le"])
            .replace("__APPELS__", str(donnees["appels_api"]))
            .replace("__DONNEES__", json.dumps(donnees, ensure_ascii=False)))
    os.makedirs(os.path.dirname(chemin), exist_ok=True)
    with open(chemin, "w", encoding="utf-8") as f:
        f.write(html)


def resume_console(donnees):
    largeur = 86
    print(f"\n{'─' * largeur}")
    print(f"  Tableau de bord — {len(donnees['etudiants'])} étudiant(s), "
          f"{len(donnees['tds'])} TD, {donnees['appels_api']} appels API")
    print(f"{'─' * largeur}")
    print(f"  {'Étudiant':<26}{'Pseudo':<20}{'Cellules':<12}{'Note':<12}{'À voir'}")
    print(f"{'─' * largeur}")
    for etudiant in donnees["etudiants"]:
        ok = sum(t["audit"]["ok"] for t in etudiant["tds"].values())
        total = sum(t["audit"]["total"] for t in etudiant["tds"].values())
        cellules = f"{ok}/{total}"
        note = f"{etudiant['total']}/{etudiant['max']}"
        statut = "✅" if etudiant["depot_ok"] else "dépôt introuvable"
        print(f"  {etudiant['nom'][:24]:<26}{(etudiant['username'] or '—')[:18]:<20}"
              f"{cellules:<12}{note:<12}{statut}")
    print(f"{'─' * largeur}\n")


def main():
    analyseur = argparse.ArgumentParser(description="Tableau de bord enseignant")
    analyseur.add_argument("td", nargs="?", help="limiter à un TD (ex. : td03)")
    analyseur.add_argument("--td", dest="td_option", help="limiter à un TD")
    analyseur.add_argument("--refresh", action="store_true", help="ignorer le cache GitHub")
    analyseur.add_argument("--offline", action="store_true", help="aucune requête réseau")
    analyseur.add_argument("--out", default=SORTIE, help="fichier HTML de sortie")
    arguments = analyseur.parse_args()

    choix = arguments.td_option or arguments.td
    tds = [choix] if choix else TDS
    for td in tds:
        if td not in TDS:
            print(f"❌ TD inconnu : {td} (attendu : {', '.join(TDS)})")
            sys.exit(1)

    if not os.path.isdir(REPONSES):
        print(f"❌ Dossier introuvable : {REPONSES}")
        print("   Déposez-y les exports Google Forms (présence, quiz).")
        sys.exit(1)

    print("📥 Lecture des exports et interrogation de GitHub…")
    donnees = collecter(tds, hors_ligne=arguments.offline, rafraichir=arguments.refresh)
    resume_console(donnees)
    ecrire_html(donnees, arguments.out)
    print(f"📊 Tableau de bord écrit : {arguments.out}")
    print(f"   Annuaire : {MAPPING}")


if __name__ == "__main__":
    main()
